import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter as L
rows=[]
for r in csv.reader(open('raw.csv')):
    name,tipo=r[0],r[1]
    key = name if tipo=="par" else name+("s" if tipo=="suited" else "o")
    combos = 6 if tipo=="par" else (4 if tipo=="suited" else 12)
    rows.append([key,{"par":"Pareja","suited":"Suited","offsuit":"Offsuit"}[tipo],combos]+[int(x) for x in r[2:]])
assert len(rows)==169 and sum(r[2] for r in rows)==1326
# validacion exacta contra conteo clasico de 7 cartas
known=[23294460,58627800,31433400,6461620,6180020,4047644,3473184,224848,41584]
for i in range(9):
    s=sum(r[2]*r[4+i] for r in rows)
    assert s%21==0 and s//21==known[i],(i,s//21,known[i])
print("VALIDACION EXACTA OK: coincide con el conteo clasico de manos de 7 cartas")
A=Font(name="Arial",size=10); B=Font(name="Arial",size=10,bold=True); BL=Font(name="Arial",size=10,color="0000FF")
H=PatternFill("solid",fgColor="1F3864"); HF=Font(name="Arial",size=10,bold=True,color="FFFFFF")
wb=Workbook()
# ---- Guia
g=wb.active; g.title="Guía"
txt=["Potencial preflop de las 169 manos iniciales de Texas Hold'em",
"",
"Qué es: para cada mano, la probabilidad de terminar (al llegar al river) con cada categoría de mano.",
"Método: enumeración EXACTA. Para cada mano se recorren los C(50,5) = 2.118.760 tableros posibles y se cuenta en cuántos sale cada categoría. Sin azar ni simulación.",
"Categoría = mejor mano de 5 cartas entre tus 2 y las 5 del tablero (7 cartas).",
"Las categorías de 'Probabilidades' son excluyentes y suman 100 %. Las columnas acumuladas ('o mejor') suman las categorías superiores.",
"'Mejora al tablero': el tablero solo (5 cartas) tiene una categoría; aquí cuenta cuando tus 2 cartas suben esa categoría. Es una medida aproximada de cuánto aportan tus cartas (no detecta mejoras dentro de la misma categoría, p. ej. un color más alto).",
"Los palos son simétricos: se calculó una mano representativa por tipo (pareja, suited, offsuit). 'Combos' = cuántas manos concretas de las 1.326 hay de ese tipo (6, 4 o 12).",
"Validación: el promedio ponderado por combos coincide EXACTO con el conteo clásico de las 133.784.560 manos de 7 cartas (hoja 'Validación').",
"",
"Ojo: potencial NO es equity. Una escalera baja pierde contra un color o una escalera mayor. Esta hoja mide versatilidad, no probabilidad de ganar.",
"Mapas: arriba de la diagonal = suited, abajo = offsuit, diagonal = parejas.",
"Hojas: Conteos (datos crudos, en azul) → Probabilidades (fórmulas) → Mapas y Validación."]
for i,t in enumerate(txt,1):
    g.cell(i,1,t).font=B if i==1 else A
    g.cell(i,1).alignment=Alignment(wrap_text=True,vertical="top")
g.column_dimensions["A"].width=110
# ---- Conteos
c=wb.create_sheet("Conteos")
cats=["Carta alta","Pareja","Doble pareja","Trío","Escalera","Color","Full house","Póker","Escalera de color"]
hdr=["Mano","Tipo","Combos","Tableros"]+cats+["Mejora al tablero"]
for j,h in enumerate(hdr,1):
    x=c.cell(1,j,h); x.font=HF; x.fill=H; x.alignment=Alignment(horizontal="center",wrap_text=True)
for i,r in enumerate(rows,2):
    for j,v in enumerate(r,1):
        x=c.cell(i,j,v); x.font=BL if j>=3 else A
        if j>=3: x.number_format="#,##0"
c.freeze_panes="B2"
for j in range(1,15): c.column_dimensions[L(j)].width=13
# ---- Probabilidades
p=wb.create_sheet("Probabilidades")
ph=["Mano","Tipo","Combos"]+cats+["Escalera o mejor","Color o mejor","Trío o mejor","Mejora al tablero"]
for j,h in enumerate(ph,1):
    x=p.cell(1,j,h); x.font=HF; x.fill=H; x.alignment=Alignment(horizontal="center",wrap_text=True)
for i in range(2,171):
    p.cell(i,1,f"=Conteos!A{i}"); p.cell(i,2,f"=Conteos!B{i}"); p.cell(i,3,f"=Conteos!C{i}")
    for k in range(9):   # E..M en Conteos -> D..L
        p.cell(i,4+k,f"=Conteos!{L(5+k)}{i}/Conteos!$D{i}")
    p.cell(i,13,f"=SUM(H{i}:L{i})"); p.cell(i,14,f"=SUM(I{i}:L{i})"); p.cell(i,15,f"=SUM(G{i}:L{i})")
    p.cell(i,16,f"=Conteos!N{i}/Conteos!D{i}")
    for j in range(1,17):
        p.cell(i,j).font=A
        if j>=4: p.cell(i,j).number_format="0.00%"
p.freeze_panes="B2"; p.auto_filter.ref="A1:P170"
for j in range(1,17): p.column_dimensions[L(j)].width=12
p.row_dimensions[1].height=30
# ---- Mapas
m=wb.create_sheet("Mapas")
R="AKQJT98765432"
def mapa(top,title,col):
    m.cell(top,1,title).font=B
    for j,ch in enumerate(R):
        for (r_,c_) in ((top+1,2+j),(top+2+j,1)):
            x=m.cell(r_,c_,ch); x.font=HF; x.fill=H; x.alignment=Alignment(horizontal="center")
    for i,a in enumerate(R):
        for j,b in enumerate(R):
            key = a+a if i==j else (a+b+"s" if i<j else b+a+"o")
            x=m.cell(top+2+i,2+j,f'=INDEX(Probabilidades!${col}$2:${col}$170,MATCH("{key}",Probabilidades!$A$2:$A$170,0))')
            x.number_format="0.0%"; x.font=Font(name="Arial",size=9); x.alignment=Alignment(horizontal="center")
    rng=f"B{top+2}:N{top+14}"
    m.conditional_formatting.add(rng,ColorScaleRule(start_type="min",start_color="F8696B",mid_type="percentile",mid_value=50,mid_color="FFEB84",end_type="max",end_color="63BE7B"))
mapa(1,"Escalera o mejor (probabilidad por el river)","M")
mapa(18,"Color o mejor (color, full, póker o escalera de color)","N")
mapa(35,"Tus cartas mejoran la categoría del tablero","P")
for j in range(1,15): m.column_dimensions[L(j)].width=8
# ---- Validacion
v=wb.create_sheet("Validación")
vh=["Categoría","Promedio ponderado (esta hoja)","Frecuencia clásica 7 cartas","Manos de 7 cartas (conteo clásico)","Diferencia"]
for j,h in enumerate(vh,1):
    x=v.cell(1,j,h); x.font=HF; x.fill=H; x.alignment=Alignment(horizontal="center",wrap_text=True)
v.cell(12,1,"Total de manos de 7 cartas").font=B; v.cell(12,4,"=SUM(D2:D10)").number_format="#,##0"
for k in range(9):
    i=2+k; col=L(4+k)
    v.cell(i,1,cats[k])
    v.cell(i,2,f"=SUMPRODUCT(Probabilidades!$C$2:$C$170,Probabilidades!{col}$2:{col}$170)/SUM(Probabilidades!$C$2:$C$170)").number_format="0.0000%"
    v.cell(i,4,known[k]).font=BL; v.cell(i,4).number_format="#,##0"
    v.cell(i,3,f"=D{i}/$D$12").number_format="0.0000%"
    v.cell(i,5,f"=B{i}-C{i}").number_format="0.000000%"
    for j in (1,2,3,5): v.cell(i,j).font=A
v.cell(14,1,"Fuente de la columna D: conteo clásico de manos de 7 cartas (133.784.560 en total; la escalera de color incluye la escalera real). El programa además comprobó la igualdad exacta con enteros.").font=A
for j,w in enumerate([28,26,24,30,16],1): v.column_dimensions[L(j)].width=w
wb.save("/mnt/user-data/outputs/potencial_preflop_169_manos.xlsx")
